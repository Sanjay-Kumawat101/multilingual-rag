"""
document_loader.py - turn an uploaded file into plain text.

Design: every supported format has ONE small loader function that returns a list
of PageText objects. To support a new format later (e.g. .pptx), write one
function and add it to the LOADERS dictionary at the bottom - nothing else changes.

Security: uploaded files are only PARSED as data. They are never executed,
imported or evaluated.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import config


# ---------------------------------------------------------------------------
# Custom exceptions (the UI catches these and shows the message directly)
# ---------------------------------------------------------------------------
class UnsupportedFormatError(Exception):
    def __init__(self, message: str = "⚠️ This file format is not currently supported."):
        super().__init__(message)


class EmptyDocumentError(Exception):
    def __init__(self, message: str = "⚠️ No readable text was found in this document."):
        super().__init__(message)


class DocumentReadError(Exception):
    """The file is corrupted, password-protected, or otherwise unreadable."""


@dataclass
class PageText:
    page: int | None   # 1-based page number, or None if the format has no pages
    text: str


# ---------------------------------------------------------------------------
# One loader per format
# ---------------------------------------------------------------------------
def _load_pdf(path: Path) -> list[PageText]:
    import pymupdf  # PyMuPDF (older tutorials use "import fitz" - same library)

    pages: list[PageText] = []
    with pymupdf.open(path) as doc:
        for number, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if text.strip():
                pages.append(PageText(page=number, text=text))
    return pages


def _load_docx(path: Path) -> list[PageText]:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Tables are not part of doc.paragraphs, so read them separately.
    # (Their original position in the document is not preserved - fine for RAG.)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    # DOCX has no fixed pages (they depend on rendering), so page = None.
    return [PageText(page=None, text="\n\n".join(parts))]


def _load_txt(path: Path) -> list[PageText]:
    raw = path.read_bytes()
    # Try the common encodings; utf-8-sig also strips a Windows BOM.
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return [PageText(page=None, text=raw.decode(encoding))]
        except UnicodeDecodeError:
            continue
    return [PageText(page=None, text=raw.decode("utf-8", errors="replace"))]


LOADERS: dict[str, Callable[[Path], list[PageText]]] = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".txt": _load_txt,
}


# ---------------------------------------------------------------------------
# Public functions used by app.py
# ---------------------------------------------------------------------------
def save_uploaded_file(filename: str, data: bytes) -> Path:
    """Save the bytes of an uploaded file into data/uploads/ and return its path."""
    safe_name = Path(filename).name          # drops folders, blocks "../../evil.py"
    if Path(safe_name).suffix.lower() not in LOADERS:
        raise UnsupportedFormatError()
    config.ensure_dirs()
    path = config.UPLOAD_DIR / safe_name
    path.write_bytes(data)
    return path


def load_document(path: str | Path) -> list[PageText]:
    """Extract text from a PDF / DOCX / TXT file. Returns non-empty pages only."""
    path = Path(path)
    loader = LOADERS.get(path.suffix.lower())
    if loader is None:
        raise UnsupportedFormatError()

    try:
        pages = loader(path)
    except Exception as exc:  # corrupted PDF, password-protected file, etc.
        raise DocumentReadError(f"⚠️ Could not read '{path.name}': {exc}") from exc

    pages = [p for p in pages if p.text.strip()]
    if not pages:  # e.g. a scanned PDF that contains only images
        raise EmptyDocumentError()
    return pages
